#!/usr/bin/env python3
"""Daily Sales Report for ERPNext - Today"""
import json, sys
import frappe
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO

def run_daily_report(today="2026-06-15"):
    # Overall stats
    overall = frappe.db.sql("""
        SELECT COUNT(*) as invoice_count, SUM(grand_total) as total_sales,
               COUNT(DISTINCT customer) as unique_customers
        FROM `tabSales Invoice`
        WHERE posting_date = %s AND docstatus = 1
    """, (today,), as_dict=True)[0]

    # By cashier
    cashiers = frappe.db.sql("""
        SELECT custom_sepidar_invoice_creator_full_name as cashier,
               COUNT(*) as invoice_count,
               SUM(grand_total) as total_sales,
               COUNT(DISTINCT customer) as customers
        FROM `tabSales Invoice`
        WHERE posting_date = %s AND docstatus = 1
          AND custom_sepidar_invoice_creator_full_name IS NOT NULL
          AND custom_sepidar_invoice_creator_full_name != ''
        GROUP BY custom_sepidar_invoice_creator_full_name
        ORDER BY total_sales DESC
    """, (today,), as_dict=True)

    # Hourly breakdown
    hourly = frappe.db.sql("""
        SELECT HOUR(posting_time) as hour,
               COUNT(*) as invoice_count,
               SUM(grand_total) as total_sales
        FROM `tabSales Invoice`
        WHERE posting_date = %s AND docstatus = 1
        GROUP BY HOUR(posting_time)
        ORDER BY hour
    """, (today,), as_dict=True)

    # By customer
    customers = frappe.db.sql("""
        SELECT customer, customer_name,
               COUNT(*) as invoice_count,
               SUM(grand_total) as total_sales
        FROM `tabSales Invoice`
        WHERE posting_date = %s AND docstatus = 1
        GROUP BY customer, customer_name
        ORDER BY total_sales DESC
        LIMIT 20
    """, (today,), as_dict=True)

    # Top items
    top_items = frappe.db.sql("""
        SELECT sii.item_code, COALESCE(i.item_name, sii.item_code) as item_name,
               SUM(sii.stock_qty) as total_qty,
               SUM(sii.amount) as total_amount
        FROM `tabSales Invoice Item` sii
        INNER JOIN `tabSales Invoice` si ON si.name = sii.parent
        LEFT JOIN `tabItem` i ON i.name = sii.item_code
        WHERE si.posting_date = %s AND si.docstatus = 1
        GROUP BY sii.item_code, COALESCE(i.item_name, sii.item_code)
        ORDER BY total_amount DESC
        LIMIT 15
    """, (today,), as_dict=True)

    # Build Excel
    wb = openpyxl.Workbook()
    hf = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    b = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
    green_fill = PatternFill(start_color="C6EFCE",end_color="C6EFCE",fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C",end_color="FFEB9C",fill_type="solid")
    gold = PatternFill(start_color="FFD700",end_color="FFD700",fill_type="solid")
    silver = PatternFill(start_color="C0C0C0",end_color="C0C0C0",fill_type="solid")
    bronze = PatternFill(start_color="CD7F32",end_color="CD7F32",fill_type="solid")

    def set_header(ws, row, headers):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = hf; cell.fill = hfill
            cell.alignment = Alignment(horizontal='center'); cell.border = b

    # Sheet 1: Summary
    ws1 = wb.active; ws1.title = "خلاصه امروز"
    ws1.merge_cells('A1:D1')
    ws1['A1'] = f"گزارش فروش امروز — {today}"
    ws1['A1'].font = Font(bold=True, size=16, color="2F5496")
    ws1['A1'].alignment = Alignment(horizontal='center')

    stats = [
        ("تعداد فاکتور", f"{overall['invoice_count']:,}"),
        ("مبلغ کل فروش (تومان)", f"{overall['total_sales']:,.0f}"),
        ("تعداد مشتری یکتا", overall['unique_customers']),
        ("میانگین فاکتور (تومان)", f"{overall['total_sales']/overall['invoice_count']:,.0f}" if overall['invoice_count'] else "0"),
    ]
    for i, (label, value) in enumerate(stats, 3):
        ws1.cell(row=i, column=1, value=label).font = Font(bold=True, size=12)
        ws1.cell(row=i, column=2, value=value)
        ws1.cell(row=i, column=2).font = Font(size=12, color="2F5496")
    ws1.column_dimensions['A'].width = 25
    ws1.column_dimensions['B'].width = 25

    # Sheet 2: Cashiers
    ws2 = wb.create_sheet("صندوق‌دارها")
    ws2.merge_cells('A1:E1')
    ws2['A1'] = "عملکرد صندوق‌دارها"
    ws2['A1'].font = Font(bold=True, size=14, color="2F5496")
    set_header(ws2, 3, ["رتبه", "نام صندوق‌دار", "تعداد فاکتور", "مبلغ فروش (تومان)", "تعداد مشتری"])
    for i, c in enumerate(cashiers, 1):
        row = i + 3
        fill = gold if i==1 else (silver if i==2 else (bronze if i==3 else None))
        ws2.cell(row=row,column=1,value=i).border=b
        ws2.cell(row=row,column=1).alignment=Alignment(horizontal='center')
        if fill: ws2.cell(row=row,column=1).fill=fill
        ws2.cell(row=row,column=2,value=c['cashier']).border=b
        if fill: ws2.cell(row=row,column=2).fill=fill
        ws2.cell(row=row,column=3,value=c['invoice_count']).border=b
        ws2.cell(row=row,column=3).alignment=Alignment(horizontal='center')
        if fill: ws2.cell(row=row,column=3).fill=fill
        ws2.cell(row=row,column=4,value=round(c['total_sales'],0)).border=b
        ws2.cell(row=row,column=4).number_format='#,##0'
        if fill: ws2.cell(row=row,column=4).fill=fill
        ws2.cell(row=row,column=5,value=c['customers']).border=b
        ws2.cell(row=row,column=5).alignment=Alignment(horizontal='center')
        if fill: ws2.cell(row=row,column=5).fill=fill
    for c,w in enumerate([8,30,15,20,15],1):
        ws2.column_dimensions[get_column_letter(c)].width=w

    # Sheet 3: Hourly
    ws3 = wb.create_sheet("تفکیک ساعتی")
    ws3.merge_cells('A1:C1')
    ws3['A1'] = "فروش به تفکیک ساعت"
    ws3['A1'].font = Font(bold=True, size=14, color="2F5496")
    set_header(ws3, 3, ["ساعت", "تعداد فاکتور", "مبلغ فروش (تومان)"])
    for i, h in enumerate(hourly, 4):
        ws3.cell(row=i,column=1,value=f"{h['hour']:02d}:00").border=b
        ws3.cell(row=i,column=1).alignment=Alignment(horizontal='center')
        ws3.cell(row=i,column=2,value=h['invoice_count']).border=b
        ws3.cell(row=i,column=2).alignment=Alignment(horizontal='center')
        ws3.cell(row=i,column=3,value=round(h['total_sales'],0)).border=b
        ws3.cell(row=i,column=3).number_format='#,##0'
    for c,w in enumerate([12,15,20],1):
        ws3.column_dimensions[get_column_letter(c)].width=w

    # Sheet 4: Top Customers
    ws4 = wb.create_sheet("مشتریان")
    ws4.merge_cells('A1:D1')
    ws4['A1'] = "بهترین مشتریان امروز"
    ws4['A1'].font = Font(bold=True, size=14, color="2F5496")
    set_header(ws4, 3, ["#", "مشتری", "تعداد فاکتور", "مبلغ خرید (تومان)"])
    for i, c in enumerate(customers, 4):
        ws4.cell(row=i,column=1,value=i-3).border=b
        ws4.cell(row=i,column=1).alignment=Alignment(horizontal='center')
        name = c['customer_name'] or c['customer']
        ws4.cell(row=i,column=2,value=name).border=b
        ws4.cell(row=i,column=3,value=c['invoice_count']).border=b
        ws4.cell(row=i,column=3).alignment=Alignment(horizontal='center')
        ws4.cell(row=i,column=4,value=round(c['total_sales'],0)).border=b
        ws4.cell(row=i,column=4).number_format='#,##0'
    for c,w in enumerate([5,30,15,20],1):
        ws4.column_dimensions[get_column_letter(c)].width=w

    # Sheet 5: Top Items
    ws5 = wb.create_sheet("کالاها")
    ws5.merge_cells('A1:D1')
    ws5['A1'] = "پرفروش‌ترین کالاها"
    ws5['A1'].font = Font(bold=True, size=14, color="2F5496")
    set_header(ws5, 3, ["#", "نام کالا", "تعداد", "مبلغ فروش (تومان)"])
    for i, item in enumerate(top_items, 4):
        ws5.cell(row=i,column=1,value=i-3).border=b
        ws5.cell(row=i,column=1).alignment=Alignment(horizontal='center')
        ws5.cell(row=i,column=2,value=item['item_name']).border=b
        ws5.cell(row=i,column=3,value=round(item['total_qty'],2)).border=b
        ws5.cell(row=i,column=3).alignment=Alignment(horizontal='center')
        ws5.cell(row=i,column=4,value=round(item['total_amount'],0)).border=b
        ws5.cell(row=i,column=4).number_format='#,##0'
    for c,w in enumerate([5,35,12,20],1):
        ws5.column_dimensions[get_column_letter(c)].width=w

    # Upload
    output = BytesIO(); wb.save(output); output.seek(0)
    file_name = f"Daily_Sales_{today}.xlsx"
    existing = frappe.db.get_value("File",{"file_name":file_name},"name")
    if existing: frappe.delete_doc("File", existing)
    doc = frappe.get_doc({"doctype":"File","file_name":file_name,"content":output.getvalue(),"is_private":0})
    doc.insert(ignore_permissions=True); frappe.db.commit()

    return {
        "success": True,
        "file_url": f"https://dehati.ir{doc.file_url}",
        "total_invoices": overall['invoice_count'],
        "total_sales": overall['total_sales'],
        "customers": overall['unique_customers']
    }

result = run_daily_report()
print(json.dumps(result))
